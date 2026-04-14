import io
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .reports import MainBalanceReport, SavingsReport, LoanReport, GlobalReport


# Shared styles
HEADER_FONT = Font(bold=True, color="FFFFFF", size=12)
HEADER_FILL = PatternFill(start_color="8B5CF6", end_color="8B5CF6", fill_type="solid")
SUBHEADER_FILL = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
MONEY_FORMAT = '#,##0.00'
THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin'),
)


def _style_header_row(ws, row, num_cols):
    """Apply header styling to a row."""
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center')
        cell.border = THIN_BORDER


def _auto_width(ws):
    """Auto-adjust column widths."""
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_length + 4, 30)


def export_weekly_balance(user, target_date=None):
    """Export weekly balance report as .xlsx"""
    report = MainBalanceReport.weekly(user, target_date)
    wb = Workbook()
    ws = wb.active
    ws.title = "Haftalik Hisobot"

    # Title
    ws.merge_cells('A1:H1')
    ws['A1'] = f"Haftalik Moliyaviy Hisobot ({report['week_start']} — {report['week_end']})"
    ws['A1'].font = Font(bold=True, size=14, color="8B5CF6")

    # Headers
    headers = ['Kun', 'Kun nomi', 'Daromad', 'Xarajat', "Jamg'armaga", "Jamg'armadan", 'Qarz olindi', 'Qarz berildi']
    for col, header in enumerate(headers, 1):
        ws.cell(row=3, column=col, value=header)
    _style_header_row(ws, 3, len(headers))

    # Data
    for i, day in enumerate(report['daily'], 4):
        ws.cell(row=i, column=1, value=day['day'])
        ws.cell(row=i, column=2, value=day['day_name'])
        ws.cell(row=i, column=3, value=day['income']).number_format = MONEY_FORMAT
        ws.cell(row=i, column=4, value=day['expense']).number_format = MONEY_FORMAT
        ws.cell(row=i, column=5, value=day['savings_deposit']).number_format = MONEY_FORMAT
        ws.cell(row=i, column=6, value=day['savings_withdraw']).number_format = MONEY_FORMAT
        ws.cell(row=i, column=7, value=day['loan_in']).number_format = MONEY_FORMAT
        ws.cell(row=i, column=8, value=day['loan_out']).number_format = MONEY_FORMAT

    # Totals row
    total_row = len(report['daily']) + 4
    ws.cell(row=total_row, column=1, value="JAMI").font = Font(bold=True)
    ws.cell(row=total_row, column=3, value=report['totals']['income']).number_format = MONEY_FORMAT
    ws.cell(row=total_row, column=4, value=report['totals']['expense']).number_format = MONEY_FORMAT
    ws.cell(row=total_row, column=3).font = Font(bold=True)
    ws.cell(row=total_row, column=4).font = Font(bold=True)

    _auto_width(ws)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer, f"haftalik_hisobot_{report['week_start']}.xlsx"


def export_monthly_all(user, year, month):
    """Export monthly report with all sections as .xlsx"""
    wb = Workbook()

    # Sheet 1: Balance
    balance_report = MainBalanceReport.monthly(user, year, month)
    ws1 = wb.active
    ws1.title = "Balans"

    ws1.merge_cells('A1:D1')
    ws1['A1'] = f"Oylik Balans Hisoboti — {year}/{month:02d}"
    ws1['A1'].font = Font(bold=True, size=14, color="8B5CF6")

    headers = ['Kun', 'Daromad', 'Xarajat', 'Farq']
    for col, h in enumerate(headers, 1):
        ws1.cell(row=3, column=col, value=h)
    _style_header_row(ws1, 3, len(headers))

    for i, day in enumerate(balance_report['daily'], 4):
        ws1.cell(row=i, column=1, value=day['day'])
        ws1.cell(row=i, column=2, value=day['income']).number_format = MONEY_FORMAT
        ws1.cell(row=i, column=3, value=day['expense']).number_format = MONEY_FORMAT
        ws1.cell(row=i, column=4, value=day['income'] - day['expense']).number_format = MONEY_FORMAT

    totals = balance_report['totals']
    t_row = len(balance_report['daily']) + 4
    ws1.cell(row=t_row, column=1, value="JAMI").font = Font(bold=True)
    ws1.cell(row=t_row, column=2, value=totals.get('income', 0)).number_format = MONEY_FORMAT
    ws1.cell(row=t_row, column=3, value=totals.get('expense', 0)).number_format = MONEY_FORMAT
    _auto_width(ws1)

    # Sheet 2: Savings
    savings_report = SavingsReport.monthly(user, year, month)
    ws2 = wb.create_sheet("Jamg'arma")

    ws2.merge_cells('A1:F1')
    ws2['A1'] = f"Oylik Jamg'arma Hisoboti — {year}/{month:02d}"
    ws2['A1'].font = Font(bold=True, size=14, color="8B5CF6")

    headers = ['Nomi', 'Hozirgi', 'Maqsad', "Qo'shildi", 'Yechildi', 'Sof']
    for col, h in enumerate(headers, 1):
        ws2.cell(row=3, column=col, value=h)
    _style_header_row(ws2, 3, len(headers))

    for i, acc in enumerate(savings_report['accounts'], 4):
        ws2.cell(row=i, column=1, value=acc['name'])
        ws2.cell(row=i, column=2, value=acc['current_amount']).number_format = MONEY_FORMAT
        ws2.cell(row=i, column=3, value=acc['target_amount']).number_format = MONEY_FORMAT
        ws2.cell(row=i, column=4, value=acc['deposits']).number_format = MONEY_FORMAT
        ws2.cell(row=i, column=5, value=acc['withdrawals']).number_format = MONEY_FORMAT
        ws2.cell(row=i, column=6, value=acc['net']).number_format = MONEY_FORMAT
    _auto_width(ws2)

    # Sheet 3: Loans
    loan_report = LoanReport.monthly(user, year, month)
    ws3 = wb.create_sheet("Qarzlar")

    ws3.merge_cells('A1:C1')
    ws3['A1'] = f"Oylik Qarz Hisoboti — {year}/{month:02d}"
    ws3['A1'].font = Font(bold=True, size=14, color="8B5CF6")

    ws3.cell(row=3, column=1, value="Ko'rsatgich")
    ws3.cell(row=3, column=2, value="Qiymat")
    _style_header_row(ws3, 3, 2)

    rows_data = [
        ("Yangi qarzlar soni", loan_report['new_loans'].get('count', 0)),
        ("Olingan qarzlar", loan_report['new_loans'].get('incoming', 0)),
        ("Berilgan qarzlar", loan_report['new_loans'].get('outgoing', 0)),
        ("To'lovlar soni", loan_report['repayments'].get('count', 0)),
        ("To'lovlar summasi", loan_report['repayments'].get('total', 0)),
    ]
    for i, (label, val) in enumerate(rows_data, 4):
        ws3.cell(row=i, column=1, value=label)
        cell = ws3.cell(row=i, column=2, value=val)
        if isinstance(val, float):
            cell.number_format = MONEY_FORMAT
    _auto_width(ws3)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer, f"oylik_hisobot_{year}_{month:02d}.xlsx"


def export_yearly_all(user, year):
    """Export yearly comprehensive report as .xlsx"""
    wb = Workbook()

    # Sheet 1: Balance yearly
    report = MainBalanceReport.yearly(user, year)
    ws1 = wb.active
    ws1.title = "Yillik Balans"

    ws1.merge_cells('A1:D1')
    ws1['A1'] = f"Yillik Moliyaviy Hisobot — {year}"
    ws1['A1'].font = Font(bold=True, size=14, color="8B5CF6")

    headers = ['Oy', 'Daromad', 'Xarajat', 'Sof']
    for col, h in enumerate(headers, 1):
        ws1.cell(row=3, column=col, value=h)
    _style_header_row(ws1, 3, len(headers))

    for i, m in enumerate(report['monthly'], 4):
        ws1.cell(row=i, column=1, value=m['month'])
        ws1.cell(row=i, column=2, value=m['income']).number_format = MONEY_FORMAT
        ws1.cell(row=i, column=3, value=m['expense']).number_format = MONEY_FORMAT
        ws1.cell(row=i, column=4, value=m['net']).number_format = MONEY_FORMAT
    _auto_width(ws1)

    # Sheet 2: Savings yearly
    savings = SavingsReport.yearly(user, year)
    ws2 = wb.create_sheet("Yillik Jamg'arma")

    ws2['A1'] = f"Yillik Jamg'arma Hisoboti — {year}"
    ws2['A1'].font = Font(bold=True, size=14, color="8B5CF6")
    row = 3
    for acc in savings['accounts']:
        ws2.cell(row=row, column=1, value=acc['name']).font = Font(bold=True, size=12)
        ws2.cell(row=row, column=2, value=f"Hozirgi: {acc['current_amount']:,.0f}")
        row += 1
        ws2.cell(row=row, column=1, value="Oy")
        ws2.cell(row=row, column=2, value="Qo'shildi")
        ws2.cell(row=row, column=3, value="Yechildi")
        _style_header_row(ws2, row, 3)
        row += 1
        for m in acc['monthly']:
            ws2.cell(row=row, column=1, value=m['month'])
            ws2.cell(row=row, column=2, value=m['deposits']).number_format = MONEY_FORMAT
            ws2.cell(row=row, column=3, value=m['withdrawals']).number_format = MONEY_FORMAT
            row += 1
        row += 1
    _auto_width(ws2)

    # Sheet 3: Category breakdown
    categories = GlobalReport.category_breakdown(user, year)
    ws3 = wb.create_sheet("Kategoriyalar")

    ws3['A1'] = f"Xarajatlar kategoriyasi — {year}"
    ws3['A1'].font = Font(bold=True, size=14, color="8B5CF6")

    headers = ['Kategoriya', 'Summa', 'Soni']
    for col, h in enumerate(headers, 1):
        ws3.cell(row=3, column=col, value=h)
    _style_header_row(ws3, 3, len(headers))

    for i, cat in enumerate(categories, 4):
        ws3.cell(row=i, column=1, value=cat['category'])
        ws3.cell(row=i, column=2, value=cat['total']).number_format = MONEY_FORMAT
        ws3.cell(row=i, column=3, value=cat['count'])
    _auto_width(ws3)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer, f"yillik_hisobot_{year}.xlsx"
